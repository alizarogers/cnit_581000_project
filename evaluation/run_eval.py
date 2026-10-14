"""
run_eval.py: run the Group 8 pun test set through the system and fill in the workbook.

Systems compared on the 30 test sentences:
  Stages 1-2   spaCy + BERT + WordNet alone (pipeline-only call, no API cost)
  Full system  Stages 1-2 evidence + Claude (Stage 3); this goes in "Pipeline says"
  Claude only  Claude without the evidence (only with --compare)

What it writes (to a new file; the input workbook is never changed):
  Test Set tab          "Pipeline says" (column J by default) = full system's call,
                        so column K and the Summary tab score themselves
  Dialog Questions tab  new "Chatbot answer" column: each question asked as a
                        follow-up after the pun has been analyzed
  Chatbot Output tab    per-sentence stage outputs, pivots, senses, scores, analyses
  System Comparison tab TP/FP/FN/TN, precision, recall, F1 for each system

Usage (from the project root, with ANTHROPIC_API_KEY set in .env):
    python -m evaluation.run_eval                       # full system + dialog questions
    python -m evaluation.run_eval --compare             # also run Claude alone
    python -m evaluation.run_eval --pipeline-only       # Stages 1-2 only, no API calls
    python -m evaluation.run_eval --only detection --limit 3
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from backend.analysis import parse_classification, parse_pivot, parse_stage_check
from backend.config import config
from backend.llm_adapter import LLMAdapter
from backend.punbot import PunBot, build_pipeline

HERE = Path(__file__).parent
DEFAULT_INPUT = HERE / "Group8_Pun_Test_Set.xlsx"
DEFAULT_OUTPUT_DIR = HERE / "results"

TEST_SHEET = "Test Set"
DIALOG_SHEET = "Dialog Questions"
OUTPUT_SHEET = "Chatbot Output"
COMPARE_SHEET = "System Comparison"
EVAL_TEMPERATURE = 0.0  # as deterministic as the SDK and model allow

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


# -- reading the workbook ---------------------------------------------------
def find_column(ws, header_prefix: str) -> Optional[int]:
    """1-based index of the first header cell starting with header_prefix (case-insensitive)."""
    for cell in ws[1]:
        if isinstance(cell.value, str) and cell.value.strip().lower().startswith(header_prefix.lower()):
            return cell.column
    return None


def read_test_set(ws) -> List[Dict]:
    id_col, sent_col = find_column(ws, "ID"), find_column(ws, "Sentence")
    label_col, target_col = find_column(ws, "Label"), find_column(ws, "Target word")
    cat_col = find_column(ws, "Category")
    rows = []
    for r in range(2, ws.max_row + 1):
        item_id, sentence = ws.cell(r, id_col).value, ws.cell(r, sent_col).value
        if not item_id or not sentence:
            continue
        rows.append({
            "row": r, "id": str(item_id).strip(), "sentence": str(sentence).strip(),
            "label": ws.cell(r, label_col).value if label_col else None,
            "target": ws.cell(r, target_col).value if target_col else None,
            "category": ws.cell(r, cat_col).value if cat_col else None,
        })
    return rows


def read_dialog_questions(ws) -> List[Dict]:
    qid_col, pid_col, q_col = find_column(ws, "QID"), find_column(ws, "Pun ID"), find_column(ws, "Question")
    rows = []
    for r in range(2, ws.max_row + 1):
        qid, pid, question = ws.cell(r, qid_col).value, ws.cell(r, pid_col).value, ws.cell(r, q_col).value
        if qid and pid and question:  # skips the blank and totals rows under the table
            rows.append({"row": r, "qid": str(qid).strip(), "pun_id": str(pid).strip(), "question": str(question).strip()})
    return rows


# -- running the systems ------------------------------------------------------
def mark(prediction: Optional[str], label: Optional[str]) -> str:
    return "??" if prediction is None else ("ok" if prediction == label else "XX")


def run_detection(items: List[Dict], pipeline, full_bot: Optional[PunBot], llm_bot: Optional[PunBot]) -> Dict[str, Dict]:
    results = {}
    for n, item in enumerate(items, 1):
        res = {k: item[k] for k in ("sentence", "label", "target", "category")}
        if pipeline is not None:
            stage = pipeline.analyze(item["sentence"])
            res["stage12"] = stage.to_dict()
            res["stage12_call"] = stage.call
        if full_bot is not None:
            out = full_bot.respond(item["sentence"], [], analyze=True, temperature=EVAL_TEMPERATURE)
            res.update(full_call=parse_classification(out["answer"]), full_pivot=parse_pivot(out["answer"]),
                       stage_check=parse_stage_check(out["answer"]), full_analysis=out["answer"],
                       full_sent=out["sent_message"])
        if llm_bot is not None:
            out = llm_bot.respond(item["sentence"], [], analyze=False, temperature=EVAL_TEMPERATURE)
            res.update(llm_call=parse_classification(out["answer"]), llm_pivot=parse_pivot(out["answer"]),
                       llm_analysis=out["answer"])
        results[item["id"]] = res

        parts = [f"[{n:>2}/{len(items)}] {item['id']}  truth={item['label']:<8}"]
        if "stage12_call" in res:
            pivot = (res["stage12"]["best"] or {}).get("text") or "-"
            parts.append(f"stages1-2={res['stage12_call']:<8}{mark(res['stage12_call'], item['label'])} ({pivot})")
        if full_bot is not None:
            parts.append(f"full={res['full_call'] or 'UNPARSED':<8}{mark(res['full_call'], item['label'])}")
        if llm_bot is not None:
            parts.append(f"claude-only={res['llm_call'] or 'UNPARSED':<8}{mark(res['llm_call'], item['label'])}")
        print("  " + "  ".join(parts))
    return results


def run_dialog(bot: PunBot, questions: List[Dict], items_by_id: Dict[str, Dict], results: Dict[str, Dict]) -> Dict[str, str]:
    answers = {}
    for n, q in enumerate(questions, 1):
        item = items_by_id.get(q["pun_id"])
        if item is None:
            answers[q["qid"]] = f"(Pun ID {q['pun_id']} not found on the Test Set tab)"
            continue
        res = results.setdefault(q["pun_id"], {})
        if "full_analysis" not in res:  # dialog-only run: analyze the pun first
            first = bot.respond(item["sentence"], [], analyze=True, temperature=EVAL_TEMPERATURE)
            res.update(full_analysis=first["answer"], full_sent=first["sent_message"])
        history = [{"role": "user", "content": res["full_sent"]},
                   {"role": "assistant", "content": res["full_analysis"]}]
        out = bot.respond(q["question"], history, analyze=False, temperature=EVAL_TEMPERATURE)
        answers[q["qid"]] = out["answer"]
        print(f"  [{n:>2}/{len(questions)}] {q['qid']} ({q['pun_id']}): {q['question'][:60]}")
    return answers


def confusion(results: Dict[str, Dict], key: str) -> Optional[Dict]:
    scored = [r for r in results.values() if r.get(key) and r.get("label")]
    if not scored:
        return None
    tp = sum(r["label"] == "Pun" and r[key] == "Pun" for r in scored)
    fp = sum(r["label"] == "Non-pun" and r[key] == "Pun" for r in scored)
    fn = sum(r["label"] == "Pun" and r[key] == "Non-pun" for r in scored)
    tn = sum(r["label"] == "Non-pun" and r[key] == "Non-pun" for r in scored)
    p = tp / (tp + fp) if tp + fp else None
    rc = tp / (tp + fn) if tp + fn else None
    f1 = 2 * p * rc / (p + rc) if p and rc else (0.0 if p is not None and rc is not None else None)
    homo = [r for r in scored if r["label"] == "Pun" and str(r.get("category", "")).startswith("Homographic")]
    hetero = [r for r in scored if r["label"] == "Pun" and str(r.get("category", "")).startswith("Heterographic")]
    rate = lambda rows: (sum(r[key] == "Pun" for r in rows) / len(rows)) if rows else None  # noqa: E731
    return {"TP": tp, "FP": fp, "FN": fn, "TN": tn, "Precision": p, "Recall": rc, "F1": f1,
            "Accuracy": (tp + tn) / len(scored), "Homographic recall": rate(homo),
            "Heterographic recall": rate(hetero), "Unparsed": len(results) - len(scored)}


# -- writing the workbook -----------------------------------------------------
def style_sheet(ws, widths: List[int], freeze: str = "C2") -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = WRAP_TOP
    ws.freeze_panes = freeze


def sense_text(s: Optional[dict]) -> str:
    if not s:
        return ""
    trig = f" [trigger: {s['trigger']}]" if s.get("trigger") else ""
    return f"{s['synset']}: {s['definition']}{trig}"


def write_results(wb, items: List[Dict], results: Dict[str, Dict], column: str, main_key: str) -> None:
    ws = wb[TEST_SHEET]
    col = column_index_from_string(column)
    for item in items:
        value = results.get(item["id"], {}).get(main_key)
        if value:
            ws.cell(item["row"], col, value)

    if OUTPUT_SHEET in wb.sheetnames:
        del wb[OUTPUT_SHEET]
    out = wb.create_sheet(OUTPUT_SHEET)
    headers = ["ID", "Sentence", "Label (ground truth)", "Target word (ground truth)",
               "Stages 1-2 call", "Stages 1-2 pivot", "Double-sense score", "Sense A (Stage 1)", "Sense B (Stage 1)",
               "Pivot role (Stage 2)", "Full system call", "Claude pivot", "Claude vs Stages 1-2",
               "Claude-only call", "Full system analysis", "Claude-only analysis"]
    out.append(headers)
    for item in items:
        r = results.get(item["id"])
        if not r:
            continue
        best = (r.get("stage12") or {}).get("best") or {}
        out.append([
            item["id"], r["sentence"], r["label"], r["target"],
            r.get("stage12_call"), best.get("text"), best.get("double_score"),
            sense_text(best.get("sense_a")), sense_text(best.get("sense_b")),
            f"{best.get('pos', '').lower()}, {best.get('dep')} of '{best.get('head')}'" if best else "",
            r.get("full_call") or ("UNPARSED" if "full_analysis" in r else None), r.get("full_pivot"),
            r.get("stage_check"), r.get("llm_call"), r.get("full_analysis"), r.get("llm_analysis"),
        ])
    style_sheet(out, [7, 44, 12, 14, 11, 14, 10, 40, 40, 22, 11, 14, 12, 11, 80, 60])


def write_comparison(wb, results: Dict[str, Dict], threshold: float) -> None:
    if COMPARE_SHEET in wb.sheetnames:
        del wb[COMPARE_SHEET]
    ws = wb.create_sheet(COMPARE_SHEET)
    systems = [("Stages 1-2 (spaCy + BERT + WordNet)", "stage12_call"),
               ("Full system (Stages 1-2 + Claude)", "full_call"),
               ("Claude only", "llm_call")]
    stats = [(name, confusion(results, key)) for name, key in systems]
    stats = [(n, s) for n, s in stats if s]
    ws.append(["Metric", *[n for n, _ in stats]])
    for metric in ["TP", "FP", "FN", "TN", "Precision", "Recall", "F1", "Accuracy",
                   "Homographic recall", "Heterographic recall", "Unparsed"]:
        ws.append([metric, *[s[metric] if s[metric] is not None else "n/a" for _, s in stats]])
    for row in ws.iter_rows(min_row=6, max_row=11, min_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    agree = [r.get("stage_check") for r in results.values() if r.get("stage_check")]
    ws.append([])
    ws.append([f"Stage 1-2 threshold: {threshold:.2f} (double-sense score needed for a Pun call)"])
    if agree:
        ws.append([f"Claude's check of Stages 1-2: " + ", ".join(f"{a} {agree.count(a)}" for a in ("Agree", "Partly agree", "Disagree"))])
    ws.append(["Counts are small (15 puns, 15 non-puns): one error moves a rate by about 7 points."])
    style_sheet(ws, [26, 32, 32, 22], freeze="B2")


def write_dialog(wb, questions: List[Dict], answers: Dict[str, str]) -> None:
    ws = wb[DIALOG_SHEET]
    col = find_column(ws, "Chatbot answer") or ws.max_column + 1
    header = ws.cell(1, col, "Chatbot answer")
    header.font, header.fill, header.alignment = HEADER_FONT, HEADER_FILL, WRAP_TOP
    ws.column_dimensions[get_column_letter(col)].width = 90
    for q in questions:
        if q["qid"] in answers:
            ws.cell(q["row"], col, answers[q["qid"]]).alignment = WRAP_TOP


def write_run_info(wb, info: List[str]) -> None:
    out = wb[OUTPUT_SHEET] if OUTPUT_SHEET in wb.sheetnames else wb.create_sheet(OUTPUT_SHEET)
    note_col = out.max_column + 2
    out.cell(1, note_col, "Run info").font = HEADER_FONT
    for i, line in enumerate(info, 2):
        out.cell(i, note_col, line)
    out.column_dimensions[get_column_letter(note_col)].width = 48


# -- main ---------------------------------------------------------------------
def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Run the pun test set through PunBot.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Test set workbook")
    parser.add_argument("--output", type=Path, help="Where to save results (default: evaluation/results/...)")
    parser.add_argument("--only", choices=["detection", "dialog"], help="Run just one part")
    parser.add_argument("--limit", type=int, help="Only the first N items of each part (quick test)")
    parser.add_argument("--column", default="J", help='Column for the full system\'s Pun/Non-pun calls (default "J")')
    parser.add_argument("--compare", action="store_true", help="Also run Claude alone, without the pipeline")
    parser.add_argument("--pipeline-only", action="store_true", help="Run Stages 1-2 only (no Claude, no API key needed)")
    parser.add_argument("--llm-only", action="store_true", help="Run without Stages 1-2 (Claude alone)")
    args = parser.parse_args(argv)

    if not args.input.exists():
        print(f"Input workbook not found: {args.input}")
        return 1
    output = args.output or DEFAULT_OUTPUT_DIR / f"{args.input.stem}_results_{datetime.now():%Y%m%d_%H%M}.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)

    pipeline = None
    if not args.llm_only:
        print("Loading Stages 1-2 (spaCy, BERT, WordNet)...")
        pipeline = build_pipeline(config, force=True)
    llm = None
    if not args.pipeline_only:
        try:
            llm = LLMAdapter()
        except EnvironmentError as e:
            print(e)
            return 1
    full_bot = PunBot(llm, pipeline) if llm else None
    llm_bot = PunBot(llm, None) if llm and (args.compare or args.llm_only) else None
    if args.llm_only:
        full_bot = None

    wb = openpyxl.load_workbook(args.input)
    items = read_test_set(wb[TEST_SHEET])
    by_id = {i["id"]: i for i in items}
    questions = read_dialog_questions(wb[DIALOG_SHEET]) if DIALOG_SHEET in wb.sheetnames else []
    if args.limit:
        items, questions = items[: args.limit], questions[: args.limit]

    print(f"Claude model: {config.CLAUDE_MODEL if llm else '(not used)'}\nInput: {args.input}")
    results, answers, ran = {}, {}, []
    main_key = "full_call" if full_bot else ("llm_call" if llm_bot else "stage12_call")

    if args.only != "dialog":
        print(f"\nDetection: {len(items)} sentences")
        results = run_detection(items, pipeline, full_bot, llm_bot)
        write_results(wb, items, results, args.column, main_key)
        write_comparison(wb, results, pipeline.threshold if pipeline else config.PUN_THRESHOLD)
        ran.append("detection")

    dialog_bot = full_bot or llm_bot
    if args.only != "detection" and questions and dialog_bot:
        print(f"\nDialog: {len(questions)} questions")
        answers = run_dialog(dialog_bot, questions, by_id, results)
        write_dialog(wb, questions, answers)
        ran.append("dialog")

    write_run_info(wb, [
        f"Run at: {datetime.now():%Y-%m-%d %H:%M}",
        f"Steps: {', '.join(ran)}",
        f"Column {args.column} holds: {main_key.replace('_call', '').replace('stage12', 'Stages 1-2').replace('full', 'full system').replace('llm', 'Claude only')}",
        f"Claude: {config.CLAUDE_MODEL if llm else 'not used'}",
        f"Parser: {config.SPACY_MODEL}" if pipeline else "Stages 1-2: not used",
        f"Encoder: {config.BERT_MODEL}" if pipeline else "",
        f"Threshold: {pipeline.threshold:.2f}" if pipeline else "",
    ])
    wb.calculation.fullCalcOnLoad = True  # make Excel recompute the scoring formulas on open
    wb.save(output)
    output.with_suffix(".json").write_text(json.dumps(
        {"claude_model": config.CLAUDE_MODEL, "bert_model": config.BERT_MODEL, "detection": results, "dialog": answers},
        indent=2, ensure_ascii=False, default=str))

    if results:
        print()
        for name, key in [("Stages 1-2", "stage12_call"), ("Full system", "full_call"), ("Claude only", "llm_call")]:
            s = confusion(results, key)
            if s:
                fmt = lambda v: f"{v:.2f}" if isinstance(v, float) else "n/a"  # noqa: E731
                print(f"{name:<12} TP={s['TP']:>2} FP={s['FP']:>2} FN={s['FN']:>2} TN={s['TN']:>2}  "
                      f"P={fmt(s['Precision'])} R={fmt(s['Recall'])} F1={fmt(s['F1'])}"
                      + (f"  unparsed={s['Unparsed']}" if s["Unparsed"] else ""))
    print(f"\nSaved: {output}\n       {output.with_suffix('.json')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
