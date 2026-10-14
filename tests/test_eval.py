"""
Tests for the classification parser and the evaluation script, with Claude mocked.
"""

import hashlib
from pathlib import Path

import openpyxl
import pytest

from backend.analysis import parse_classification, parse_pivot
from evaluation import run_eval

TEST_SET = Path(__file__).parent.parent / "evaluation" / "Group8_Pun_Test_Set.xlsx"


@pytest.mark.parametrize("text,expected", [
    ("**Classification**: Pun", "Pun"),
    ("1. **Classification**: Non-pun\n2. **Pivot**: bank", "Non-pun"),
    ("**Classification:** Non-Pun", "Non-pun"),
    ("Classification: not a pun", "Non-pun"),
    ("**Classification**: **Pun**", "Pun"),
    ("No label here", None),
])
def test_parse_classification(text, expected):
    assert parse_classification(text) == expected


@pytest.mark.parametrize("text,expected", [
    ("**Pivot**: interest (homographic)", "interest"),
    ("## Pivot\n**interest** (homographic)", "interest"),
    ('2. **Pivot**: "two tired", heterographic', "two tired"),
    ("- **Pivot:** hole in one (phrase-level)", "hole in one"),
    ("**Pivot**: None. No double meaning.", None),
    ("there is no pivot here", None),
])
def test_parse_pivot(text, expected):
    assert parse_pivot(text) == expected


class FakeLLM:
    """Calls every sentence containing 'interest' or 'cast' a pun, everything else a non-pun."""

    def __init__(self):
        self.calls = []

    def generate(self, system_prompt, history, user_message, temperature=None):
        self.calls.append((history, user_message, temperature))
        if history:
            return f"Follow-up answer to: {user_message}"
        is_pun = any(w in user_message for w in ("interest", "cast"))
        return f"**Classification**: {'Pun' if is_pun else 'Non-pun'}\n**Pivot**: something (test)"


def test_run_eval_fills_copy_and_keeps_original(tmp_path, monkeypatch):
    fake = FakeLLM()
    monkeypatch.setattr(run_eval, "LLMAdapter", lambda: fake)
    before = hashlib.md5(TEST_SET.read_bytes()).hexdigest()
    out = tmp_path / "results.xlsx"

    assert run_eval.main(["--input", str(TEST_SET), "--output", str(out)]) == 0

    assert hashlib.md5(TEST_SET.read_bytes()).hexdigest() == before  # input untouched
    wb = openpyxl.load_workbook(out)
    ts = wb["Test Set"]
    assert ts["A2"].value == "P01" and ts["J2"].value == "Pun"        # banker / interest
    assert ts["A3"].value == "P02" and ts["J3"].value == "Pun"        # cast
    assert ts["J4"].value == "Non-pun"                                 # velcro, missed by fake
    assert str(ts["K2"].value).startswith("=IF(")                      # scoring formula intact
    assert all(ts.cell(r, 10).value in ("Pun", "Non-pun") for r in range(2, 32))

    dq = wb["Dialog Questions"]
    ans_col = run_eval.find_column(dq, "Chatbot answer")
    assert ans_col is not None
    assert dq.cell(2, ans_col).value == "Follow-up answer to: Why is this funny?"
    assert dq.cell(15, ans_col).value is None                          # totals rows left alone
    assert "Chatbot Output" in wb.sheetnames

    # Dialog questions are asked with the pun and its analysis as history
    follow_ups = [c for c in fake.calls if c[0]]
    assert len(follow_ups) == 12
    assert follow_ups[0][0][0]["content"] == "I used to be a banker, but I lost interest."
    assert all(c[2] == run_eval.EVAL_TEMPERATURE for c in fake.calls)
    assert out.with_suffix(".json").exists()


def test_custom_column_leaves_j_empty(tmp_path, monkeypatch):
    monkeypatch.setattr(run_eval, "LLMAdapter", lambda: FakeLLM())
    out = tmp_path / "r.xlsx"
    run_eval.main(["--input", str(TEST_SET), "--output", str(out), "--only", "detection", "--column", "M"])
    ts = openpyxl.load_workbook(out)["Test Set"]
    assert ts["J2"].value is None and ts["M2"].value == "Pun"
